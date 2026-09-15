from collections import Counter
from datetime import date
from pathlib import Path
import json
import math
import sys
import unicodedata

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parent.parent
SOURCE = Path('C:/Users/Administrator/Desktop/VM-R&D- VM REC - VM FEEDBACK - 04092026.xlsx')
DATA = json.loads((ROOT / '.analysis/pm_analysis.json').read_text(encoding='utf-8'))
OUTPUT = ROOT / 'VM_REC_产品经理行动清单_20260910.xlsx'
REPORT = ROOT / 'VM_REC_产品经理分析_20260910.md'
source_book = load_workbook(SOURCE, data_only=False)
source = source_book['6-REPORT']
actual_rows = [r for r in range(16, source.max_row + 1) if source.cell(r, 5).value is not None]
assert sorted(record['row'] for record in DATA['records']) == actual_rows
assert len(actual_rows) == 39
status_counts = Counter(source.cell(r, 9).value for r in actual_rows)
version_counts = Counter(source.cell(r, 15).value for r in actual_rows)

GREEN = '19594D'
INK = '202C2A'
MUTED = '586763'
WHITE = 'FFFFFF'
LINE = 'DDE5E2'
YELLOW = 'FFF0C2'
RED = 'FBE3DE'
PALE = 'EDF5F1'
wb = Workbook()
wb.remove(wb.active)
wb.properties.title = 'VM REC 产品经理行动清单'
wb.properties.subject = '依据客户多轮反馈整理的建议行动，非正式排期或验收结果'
wb.properties.creator = 'Codex'


def display_width(text):
    return sum(2 if unicodedata.east_asian_width(ch) in ('W', 'F') else 1 for ch in str(text))


def row_height(values, widths, minimum=40, maximum=245):
    lines = 1
    for value, width in zip(values, widths):
        needed = sum(max(1, math.ceil(display_width(part) / max(width - 3, 8))) for part in str(value or '').split('\n'))
        lines = max(lines, needed)
    return min(maximum, max(minimum, lines * 14 + 10))


def setup_sheet(name, title, note, headers, rows, widths, table_name, freeze='D5'):
    ws = wb.create_sheet(name)
    last = len(headers)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last)
    ws.cell(1, 1, title).font = Font(name='Microsoft YaHei', size=18, bold=True, color=WHITE)
    ws.cell(1, 1).fill = PatternFill('solid', fgColor=GREEN)
    ws.cell(1, 1).alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 38
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last)
    ws.cell(2, 1, note).font = Font(name='Microsoft YaHei', size=10, color=MUTED)
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[2].height = 34
    ws.row_dimensions[3].height = 10
    for col, (header, width) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(4, col, header)
        cell.font = Font(name='Microsoft YaHei', size=10, bold=True, color=WHITE)
        cell.fill = PatternFill('solid', fgColor=GREEN)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[4].height = 32
    for r, values in enumerate(rows, 5):
        for c, value in enumerate(values, 1):
            cell = ws.cell(r, c, value)
            cell.font = Font(name='Microsoft YaHei', size=10, color=INK)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cell.border = Border(bottom=Side(style='hair', color=LINE))
        ws.row_dimensions[r].height = row_height(values, widths)
    if rows:
        tab = Table(displayName=table_name, ref=f'A4:{get_column_letter(last)}{4 + len(rows)}')
        tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium4', showRowStripes=True)
        ws.add_table(tab)
    ws.freeze_panes = freeze
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:4'
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = '第 &P 页 / 共 &N 页'
    return ws


overview = [
    ['建议结论', '先冻结 4.0.0 范围，补齐蓝牙配对、进入设置录像规则和三路时间轴等关键规格，再统一 App/固件/平台的排期与验收。客户多次选择复用已有方案以加快交付。'],
    ['分析依据', f'源文件：{SOURCE.name}\n工作表：6-REPORT；表头版本 3.9.0；本分析日期 2026-09-10。'],
    ['反馈时间顺序', 'L 列：客户 2026-08-24；N 列：凌度 2026-08-28；P 列：客户 2026-09-03。P 为空时回看 L/E 及研发答复。文件名日期不替代反馈列日期。'],
    ['有效记录', '39 条有需求正文的记录。No.1~35 含一个重复 No.7（R23 与 R42）及 3 条未编号（R17、R50、R54）；R31 为合并续行，不额外计数。'],
    ['原始状态统计', '31 Process、7 Done、1 Pending。它们是源表标签，不能据此计算真实开发完成率。Done 包括已答复、客户放弃、采用方案等不同情形。'],
    ['原始计划统计', '25 条标 4.0.0、2 条标 4.1.0、5 条标 /、7 条未填。这是行记录数，不是去重后的独立功能数，且包含已取消和跨端事项。'],
    ['原始分类统计', '24 Requirement、14 Question、1 未分类；Bug 标签为 0 不代表没有缺陷。Low 8、Medium 14、High 12、未填 5。C 列 Items 未填，顶部软硬件计数不能用于判断实际职责。'],
    ['分析与承诺边界', '本文件的行动、角色、P0/P1、建议版本和验收口径是产品建议，不是客户已批准的新范围，也不表示研发已交付。建议须经相关负责人确认。'],
    ['P0 / P1 定义', 'P0：本轮优先解决，影响范围冻结、关键流程或跨端交付；P1：同步推进的交付完善与后续范围确认。不是线上故障等级，也不覆盖源表客户优先级。'],
    ['已确认取消/降级', '取消平板/车机通用适配、自定义三导航、列表通道筛选、未下载设备视频地图、横转竖视频。启动图选静态；电子狗页面保留旧版。列表筛选取消不等于三路时间轴取消。'],
    ['最关键歧义', 'R33 进入设置原翻译不准；R38 三路切换还是同步播放未明确；R47 Point 3 在原文是变速、研发回复却是配乐；AR/SR 状态与版本冲突；N29 图片保留旧导航。'],
    ['跨端依赖', 'R41 七项默认值/隐藏由固件负责；OTA 涉及 App/平台/固件；蓝牙、三路通道、车速、AR/SR 必须按机型核验。不得凭界面图推断全部机型具备能力。'],
    ['资料覆盖', '已读取全部非空单元格并查看 50 张已挂载图片（含正文右侧 Q:T 列图片）。源文件只有一个工作表；提到的 Video 1/2 与 VM REC vs 70mai 比较工作表不在文件内。'],
    ['尚未核验', '未打开外部 Figma/YouTube 链接，未运行 App 或验证固件协议、授权价格、性能和已交付版本。本文不据此宣称实现完成。'],
    ['追溯方式', '39 条反馈结论使用 R行号作唯一标识，可点击首列跳到原始反馈。原始反馈页摘录源表文本及原行号，不含嵌入图片；图片判断在依据栏标明。'],
    ['如何推进', '先处理 A01~A06；同步推进 A07~A09；范围和依赖明确后用 A10 形成提测与发布计划。行动页的负责人和目标日期待评审填写，当前状态默认为待启动。']
]
setup_sheet('阅读说明', 'VIETMAP REC | 产品经理分析', '以最新反馈收敛范围；保留原表事实，单独标注产品建议。', ['项目', '结论与边界'], overview, [25, 128], 'Overview', freeze='B5')

task_rows = []
for task in DATA['tasks']:
    task_rows.append([
        task['id'], task['priority'], task['title'], task['action'], task['output'] + '\n完成判定：' + task['finish'],
        task['owner'], task['when'], task['refs'], '待启动', None, None
    ])
task_ws = setup_sheet(
    '你的行动清单', '10 项产品经理行动', '优先级、协作角色和时点为建议；负责人、目标日期待评审填入。',
    ['任务ID', '建议顺序', '你要推进的事', '具体行动', '产出与完成判定', '建议牵头/协作', '建议时点', '原表依据', '推进状态', '负责人', '目标日期'],
    task_rows, [10, 12, 28, 58, 60, 34, 22, 33, 16, 16, 17], 'ActionList'
)
for r, task in enumerate(DATA['tasks'], 5):
    task_ws.cell(r, 2).fill = PatternFill('solid', fgColor=RED if task['priority'] == 'P0' else YELLOW)
    task_ws.cell(r, 2).font = Font(name='Microsoft YaHei', size=10, bold=True, color=INK)
    for c in (9, 10, 11):
        task_ws.cell(r, c).fill = PatternFill('solid', fgColor=YELLOW)
    task_ws.cell(r, 11).number_format = 'yyyy-mm-dd'
dv = DataValidation(type='list', formula1='"待启动,进行中,待客户,待研发,待素材,待商务,待验收,已关闭"')
task_ws.add_data_validation(dv)
dv.add(f'I5:I{4 + len(task_rows)}')

record_rows = []
for item in DATA['records']:
    r = item['row']
    number = source.cell(r, 2).value
    ref = f'R{r} / No.{number}' if number is not None else f'R{r} / 未编号'
    record_rows.append([
        ref, item['title'], source.cell(r, 8).value or '未填', source.cell(r, 9).value or '未填',
        source.cell(r, 15).value or '未填', item['state'], item['scope'], item['action'], item['owner'], item['check'], item['evidence']
    ])
record_ws = setup_sheet(
    '39条反馈结论', '逐条归一 | 原表状态与当前判断分开', 'R31 为合并续行；R23/R42 是不同需求但都标 No.7。点击来源可跳转到本工作簿原始反馈页。',
    ['来源（点击）', '主题', '原客户优先级', '原表状态', '原表App计划', '按反馈归纳的当前状态', '建议处理/版本', '下一步动作', '建议负责角色', '验证或关闭条件', '关键证据单元格'],
    record_rows, [22, 30, 15, 14, 16, 31, 38, 64, 32, 61, 48], 'NormalizedFeedback', freeze='C5'
)
for r, item in enumerate(DATA['records'], 5):
    cell = record_ws.cell(r, 1)
    cell.hyperlink = f"#'原始反馈'!E{item['row']}"
    cell.font = Font(name='Microsoft YaHei', size=10, color=GREEN, underline='single')
    if '取消' in item['state']:
        record_ws.cell(r, 6).fill = PatternFill('solid', fgColor='E4E8E6')
    elif any(term in item['state'] for term in ('待确认', '歧义', '客户明确等待', '未解决', '翻译', '等待')):
        record_ws.cell(r, 6).fill = PatternFill('solid', fgColor=YELLOW)

setup_sheet(
    '待确认与资料', '12 项需要收口的信息', '用于下一次对客/内部评审；这是问题清单，没有发送给任何外部联系人。',
    ['ID', '建议对接方', '确认主题', '要确认或准备什么', '为什么影响决策', '原表依据'],
    DATA['questions'], [10, 31, 29, 82, 65, 33], 'OpenQuestions'
)
setup_sheet(
    '建议验收口径', '15 组交付验收口径', '以下为建议验收方向，需在范围与机型能力确认后细化成测试用例；不是已通过的测试结果。',
    ['模块', '建议验证内容', '边界/前提', '原表依据'],
    DATA['acceptance'], [28, 99, 68, 32], 'AcceptanceGuide', freeze='B5'
)

raw = wb.create_sheet('原始反馈')
raw.sheet_view.showGridLines = False
raw.sheet_view.zoomScale = 75
raw.merge_cells('A1:P1')
raw['A1'] = '原始反馈文本摘录 | 保持源表单元格行列坐标'
raw['A1'].font = Font(name='Microsoft YaHei', size=18, bold=True, color=WHITE)
raw['A1'].fill = PatternFill('solid', fgColor=GREEN)
raw.row_dimensions[1].height = 38
raw.merge_cells('A2:P3')
raw['A2'] = f'来源：{SOURCE}\n仅摘录第 15~55 行的原始文本，不含图片；完整原件未修改。P 列为 2026-09-03 客户反馈。'
raw['A2'].font = Font(name='Microsoft YaHei', size=10, color=MUTED)
raw['A2'].alignment = Alignment(wrap_text=True, vertical='center')
raw['A5'] = '原文件'
raw['B5'] = SOURCE.name
raw['B5'].hyperlink = str(SOURCE)
raw['B5'].font = Font(name='Microsoft YaHei', size=10, underline='single', color=GREEN)
raw.merge_cells('B5:H5')
raw['A7'] = '表头版本'
raw['B7'] = str(source['E1'].value)
raw.merge_cells('B7:E7')
for r in range(8, 15):
    raw.row_dimensions[r].hidden = True
raw_widths = [15, 12, 15, 18, 85, 85, 36, 16, 16, 85, 85, 75, 75, 85, 18, 85]
for c, width in enumerate(raw_widths, 1):
    raw.column_dimensions[get_column_letter(c)].width = width
for row in source.iter_rows(min_row=15, max_row=55, max_col=16):
    values = []
    for original in row:
        cell = raw.cell(original.row, original.column, original.value)
        if isinstance(original.value, str):
            cell.data_type = 's'
        cell.font = Font(name='Microsoft YaHei', size=10, color=INK)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        cell.fill = PatternFill('solid', fgColor=WHITE if original.row % 2 else PALE)
        values.append(original.value)
    raw.row_dimensions[row[0].row].height = row_height(values, raw_widths, minimum=42, maximum=405)
for c in raw[15]:
    c.font = Font(name='Microsoft YaHei', size=10, bold=True, color=WHITE)
    c.fill = PatternFill('solid', fgColor=GREEN)
raw.row_dimensions[15].height = 45
raw.freeze_panes = 'E16'
raw.auto_filter.ref = 'A15:P55'
raw.print_title_rows = '15:15'
raw.sheet_properties.tabColor = MUTED
task_ws.sheet_properties.tabColor = GREEN
record_ws.sheet_properties.tabColor = 'C98523'
wb.active = wb.sheetnames.index('你的行动清单')
wb.save(OUTPUT)

task_table = '\n'.join(
    f"| {task['id']} / {task['priority']} | {task['title']} | {task['output']} |"
    for task in DATA['tasks']
)
record_table = '\n'.join(
    f"| R{item['row']} / {('No.' + str(source.cell(item['row'], 2).value)) if source.cell(item['row'], 2).value is not None else '未编号'} | {item['title']} | {item['state']} | {item['scope']} |"
    for item in DATA['records']
)
report = f'''# VM REC：产品经理应该推进什么

分析日期：2026-09-10。依据源文件 `VM-R&D- VM REC - VM FEEDBACK - 04092026.xlsx` 的 `6-REPORT` 页；最新客户反馈为 P 列 2026-09-03。下文 `R25` 等均表示原表行号。

**建议：先冻结 4.0.0 的交付范围，优先补蓝牙配对、进入设置时的录像规则、三路时间轴及 OTA 等规格，再拉齐固件、App、平台、商务和测试。** 客户反复选择静态图、旧电子狗页面和既有 UI，说明加快交付是本轮的明显诉求。

## 1. 先修正对表格的理解

- 实际是 **39 条反馈记录**，不是 35 个独立需求：No.7 重复（R23 登录/注册，R42 设备视频回放），R17/R50/R54 未编号；R31 只是合并续行。
- 原始状态是 **31 Process、7 Done、1 Pending**。Done 混有已答复、放弃定制、采用方案等含义，不能作为已开发或已验收证据。
- O 列原计划 **25 条 4.0.0、2 条 4.1.0、5 条 /、7 条未填**。这些是行记录数，包含重复、已取消和固件事项，不是最终版本功能数量。
- Request 为 24 Requirement、14 Question、1 未分类；Bug 标签为 0 不代表产品没有缺陷。Items 列未填，顶部软硬件数量无法反映实际跨端工作。
- 所有新增优先级、负责人角色、建议版本和验收口径都是本次分析建议；实际负责人、工期与完成情况仍需评审确认。

## 2. 你的行动顺序

P0 表示本轮优先解决的范围、交互或依赖问题，不代表线上故障等级。

| 建议顺序 | 要推进的事 | 你应拿到的产出 |
|---|---|---|
{task_table}

## 3. 最需要你亲自收口的决策

**蓝牙原型（No.9 / R25）**：客户明确写了“Wait for your sharing”，N25 明确需要凌度 App 产品经理设计原型。你应直接输出发现设备、选择设备、连接 Wi-Fi、连接结果、权限/失败/手动兜底等流程，并先核清机型及系统能力。

**进入设置的录像规则（No.16 / R33）**：客户纠正了翻译。他们要么继续录像进入设置，要么收到提示后由设备自动停录，重点是不让用户先手动点停止。建议按固件能力选方案，明确取消、停录失败、返回及恢复录像，不能继续只写“确认后进入设置”。

**抓拍入口（No.15 / R32）**：客户要移除或隐藏拍照功能入口。旧英文回复容易被理解成不支持删除照片，需要按真正诉求重新评估并正面答复，当前并未关闭。

**三路时间轴（No.21 / R38）**：客户强调前/后/车内三路竞争力；列表筛选（No.17 / R34）已取消，但时间轴仍明确要求。建议先按单路切换设计，再确认是否需要三画面同步播放。N38 的 20h 是现有范围初估，不能直接覆盖更复杂的新解释。右侧 Q38:T38 有三路参考截图。

**OTA 与设备卡片（No.11~13 / R27:R29）**：自动发现版本、卡片提醒、下载、连接设备传包、安装和结果反馈应是一条完整链路。N29 图内还写到了 Wi-Fi/互联网切换导致提醒可能无法展示，必须补实际网络策略。No.12 的 5h 不是整个 OTA 项目估时。

**AR/SR 与编辑（No.6/29/30 / R22、R47:R48）**：同一相关能力出现 Done、4.0.0、4.1.0。客户仍在等报价，需协调商务准备给其决策人确认；原视频编辑第 3 点是“变速”，N47 第 3 点却写“音乐”，所以 P47 的“Skip it”究竟取消谁必须按功能名核清。基础编辑、授权 AR/SR、AI 应分别管理；Q4 只是旧回复中的暂定评估。

**离线地图（No.26 / R44）**：客户要测试包及时间，当前只有解释。提供可测试的 App 名称、平台、版本、日期、范围和路径。No.18 设备相册地图取消，不意味着本地或实时预览的所有地图场景都取消；需先定边界再交测试包。

## 4. 建议冻结的版本边界

| 处理方式 | 内容 | 原表依据 |
|---|---|---|
| 4.0.0 以复用和必要改造为主 | 静态启动图、启动/连接优化、无账号基础 UI、连接流程、OTA、设备卡片、旧电子狗页面、时间轴、本地相册、车速、60 秒分享和通用设置。仍需补原型、能力及验收定义。 | R17:R18、R23:R30、R33、R36:R40、R43:R46、R51:R55 |
| 4.0.0 有条件项 | 三路时间轴先确认观看方式；预览首帧先验证优化空间；离线地图先明确可测试场景。 | R18、R38、R44 |
| 固件独立交付并联调 | 七项默认值与菜单隐藏；App 版本号不能代替固件计划。 | R41 |
| 原计划 4.1.0，需继续收口 | 基础编辑：先确认裁剪/配乐/变速取舍，AR 单独授权。 | R47 |
| 待商务后定版本 | AR/SR 的授权、覆盖机型和实现范围。 | R22、R47:R48 |
| 暂不承诺上线日期 | AI 编辑：后续单独评估。 | J47/K47/L47 |
| 已取消或移出原计划 | 平板/车机适配、自定义三导航、列表摄像头筛选、设备相册地图、横转竖视频。 | R16、R19、R34、R35、R49 |
| 简化方案本期采用 | 启动图选静态，不做启动动画；电子狗页保留现版，新图文/搜索/全量下载放后续。 | P17、P30 |

“移除登录”“隐藏抓拍”等是功能改动诉求，不应被误标为整个需求取消。取消设备相册地图也不应顺带取消本地速度显示。

## 5. 应拉齐的跨端规则与验收

- **固件七项设置**：曝光 0.0、日期格式 DD/MM/YYYY、时区 GMT+7、速度 km/h、Logo ON、型号 ON、日期时间 ON；固定值与菜单隐藏都需真实生效。补机型、固件版本、恢复默认及升级后行为。依据 E41/F41/N41/P41。
- **机型能力表**：R29 列出 M2、M1、TS-H3K、S720、S860、L110；蓝牙、OTA、三路索引、GPS/速度及 AR/SR 不能仅凭该列表视为全支持，需逐项核验。
- **相册映射**：N37 已给出 Normal=循环录像+未锁定延时，Event=锁存+移动侦测，Parking=停车碰撞。UI 中普通与延时分开，但文件同属 Normal，需协议字段区分；与早期 Video/Lock/Photo 方案合并确认。
- **性能**：分别记录启动可操作、连接成功、预览首帧；统一测试条件，按基线确定目标。表内没有可直接沿用的数值目标。
- **分享**：60 秒是已讨论方向；明确 App 交出的文件分辨率与体积策略。Facebook 等平台可能再处理，不承诺平台最终画质完全不降。
- **UI 基线**：R19 已取消旧定制三导航，R24 采用设备/通用基础结构，但 N29 卡片图片仍有三个底部导航，需统一。卡片图上的 Device Info 文案与三点菜单入口也应对齐。
- **素材**：先提供静态启动图规格；核验 OSD 文件是否已交付及越南语版本；R55 英文仅说条款政策不变，旧中文仍包含保修入口替换为 VML 下载链接，需核清最终范围。

不要把零散估时相加当作项目周期：N28 的 5h 为新增版本提醒初估；N30 选择旧页 2h，未选的新页 20h 不计入本期；N38 的 20h 需核清范围；N49 的 16h 转码已被取消。各端开发、平台接口、固件、联调、测试与客户验收仍需独立排期。

## 6. 逐条状态索引

详细动作、角色、证据单元格、验收建议及可填写的跟踪列见配套 Excel。

| 原表定位 | 主题 | 按反馈归纳的状态 | 建议处理/版本 |
|---|---|---|---|
{record_table}

## 7. 分析范围

已读取全部非空单元格，并查看 50 张已挂载图片，包括正文右侧 Q:T 的图片。原文件仅有 `6-REPORT` 一个工作表；提到的 Video 1、Video 2、`VM REC vs 70mai` 比较表未包含在该文件内。未打开外部 Figma/YouTube，也未实际测试 App、固件、价格或性能。因此，本文能判断反馈和范围冲突，不能证明功能已实现、已验收或某项技术方案对所有机型可行。

表内“请提供”“发费用给老板”等是客户在文档中的诉求，本次仅整理为待办，没有向任何人发送消息，也没有修改原 Excel。
'''
REPORT.write_text(report, encoding='utf-8')

saved = load_workbook(OUTPUT, data_only=False)
assert saved.sheetnames == ['阅读说明', '你的行动清单', '39条反馈结论', '待确认与资料', '建议验收口径', '原始反馈']
assert saved['你的行动清单'].max_row == 14
assert saved['39条反馈结论'].max_row == 43
assert saved['待确认与资料'].max_row == 16
assert saved['建议验收口径'].max_row == 19
for r in range(15, 56):
    for c in range(1, 17):
        assert saved['原始反馈'].cell(r, c).value == source.cell(r, c).value, (r, c)
assert all(saved['39条反馈结论'].cell(i, 1).hyperlink for i in range(5, 44))
assert len(saved['你的行动清单'].data_validations.dataValidation) == 1
assert not any(cell.data_type == 'f' for sheet in saved for row in sheet for cell in row)
print(json.dumps({
    'workbook': str(OUTPUT), 'report': str(REPORT),
    'validated': {'feedback_records': len(actual_rows), 'tasks': len(DATA['tasks']), 'questions': len(DATA['questions']), 'acceptance_groups': len(DATA['acceptance']), 'raw_cells_match_source': True},
    'source_status': dict(status_counts), 'source_version': {str(k): v for k, v in version_counts.items()}
}, ensure_ascii=False))
